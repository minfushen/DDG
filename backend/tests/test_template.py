"""模板解析与配置（非功能需求①）测试。"""
import shutil
import sys
import tempfile
from pathlib import Path

import pytest

# 让 backend 成为可导入根
BACKEND = Path(__file__).resolve().parent.parent
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.template.models import BlockType, ReportTemplate  # noqa: E402
from app.template.parser import parse_markdown, parse_template_file  # noqa: E402
from app.template.renderer import render_report_from_template  # noqa: E402
from app.template.store import (  # noqa: E402
    _RAW_DIR,
    _STORE_DIR,
    delete_template,
    get_active_template,
    get_template,
    list_templates,
    save_template,
    set_active,
)
from app.agents.sub_agents.full_report_builder import build_full_due_diligence_report  # noqa: E402
from app.agents.research_engine.synthesizer import synthesize_research_report  # noqa: E402


SAMPLE_MD = """# 企业概况
请在此填写企业基本面的静态指引。
{{解读:工商治理解读@business}}

# 财务分析
| 指标 | 单位 |
| --- | --- |
| 营业收入 | 万元 |
| 净利润 | 万元 |
{{子报告:financial}}
{{解读:财务健康度@financial}}

# 行业环境
{{子报告:industry}}
"""


@pytest.fixture()
def tmp_store(tmp_path, monkeypatch):
    store = tmp_path / "templates"
    raw = tmp_path / "raw"
    monkeypatch.setattr("app.template.store._STORE_DIR", store)
    monkeypatch.setattr("app.template.store._RAW_DIR", raw)
    store.mkdir(parents=True, exist_ok=True)
    raw.mkdir(parents=True, exist_ok=True)
    yield store


def _count_blocks(tpl, btype):
    return sum(1 for s in tpl.sections for b in s.blocks if b.type == btype)


def test_parse_markdown_structure():
    tpl = parse_markdown(SAMPLE_MD, name="测试模板")
    assert [s.title for s in tpl.sections] == ["企业概况", "财务分析", "行业环境"]
    # 指标：2 个表内 + 0 个行内 = 2
    assert _count_blocks(tpl, BlockType.INDICATOR) == 2
    assert _count_blocks(tpl, BlockType.SUBREPORT) == 2
    assert _count_blocks(tpl, BlockType.INTERPRETATION) == 2
    assert _count_blocks(tpl, BlockType.NARRATIVE) == 1
    # 指标表单位解析
    inds = [b for s in tpl.sections for b in s.blocks if b.type == BlockType.INDICATOR]
    assert {b.unit for b in inds} == {"万元"}


def test_parse_inline_slot_with_dim_and_unit():
    md = "{{指标:资产负债率|单位:%}}\n{{解读:司法合规@legal}}"
    tpl = parse_markdown(md, name="x")
    ind = [b for s in tpl.sections for b in s.blocks if b.type == BlockType.INDICATOR][0]
    assert ind.key == "资产负债率" and ind.unit == "%"
    interp = [b for s in tpl.sections for b in s.blocks if b.type == BlockType.INTERPRETATION][0]
    assert interp.source_dimension == "legal"


def test_parse_docx_roundtrip(tmp_path):
    # 构造伪 docx：无 python-docx 时跳过，避免硬依赖
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        pytest.skip("python-docx 未安装，跳过 docx 解析测试")
    p = tmp_path / "t.docx"
    doc = Document()
    doc.add_heading("章节一", level=1)
    doc.add_paragraph("静态指引文本")
    doc.add_paragraph("{{子报告:financial}}")
    doc.save(str(p))
    tpl = parse_template_file(p, name="docx模板")
    assert tpl.sections[0].title == "章节一"
    assert any(b.type == BlockType.SUBREPORT for s in tpl.sections for b in s.blocks)


def test_store_crud_and_activate(tmp_store):
    tpl = parse_markdown(SAMPLE_MD, name="我的模板")
    saved = save_template(tpl)
    assert saved.id and get_template(saved.id) is not None
    # 激活
    set_active(saved.id)
    assert get_active_template().id == saved.id
    # 列表包含内置 + 自建
    ids = [t.id for t in list_templates()]
    assert saved.id in ids and "builtin_standard" in ids
    # 删除
    assert delete_template(saved.id) is True
    assert get_template(saved.id) is None


def test_renderer_binds_subreports_indicators_interpretation():
    tpl = parse_markdown(SAMPLE_MD, name="渲染测试")
    sub_reports = {
        "financial": {
            "recommendation": "财务指标整体健康",
            "report_chapters": [{"title": "财务专项", "analysis": ["营收同比改善"]}],
        },
        "business": {"recommendation": "治理规范"},
        "industry": {"recommendation": "行业景气度中性", "report_chapters": [{"title": "行业", "analysis": ["平稳"]}]},
    }
    indicators = {"营业收入": "1200万元", "净利润": "300万元"}
    sections = render_report_from_template(tpl, sub_reports, indicators=indicators)
    assert len(sections) == 3
    # 指标绑定
    fin_sec = sections[1]
    ind_block = next(c for c in fin_sec["content"] if c["type"] == "indicator")
    assert ind_block["value"] == "1200万元"
    # 子报告绑定
    sub_block = next(c for c in fin_sec["content"] if c["type"] == "subreport" and c["dimension"] == "financial")
    assert sub_block["chapters"][0]["title"] == "财务专项"
    # 解读绑定（取 financial.recommendation）
    interp_block = next(c for c in fin_sec["content"] if c["type"] == "interpretation")
    assert interp_block["text"] == "财务指标整体健康"


def test_renderer_indicator_fallback_when_missing():
    tpl = parse_markdown("# 财务\n{{指标:营业收入|单位:万元}}", name="x")
    sections = render_report_from_template(tpl, {}, indicators=None)
    assert sections[0]["content"][0]["value"] == "待补充"


def test_full_report_builder_attaches_template_sections():
    tpl = parse_markdown(SAMPLE_MD, name="集成测试")
    sub_reports = {
        "business": {"recommendation": "治理规范"},
        "financial": {
            "recommendation": "财务健康",
            "key_metrics": {"revenue": "1200", "net_profit": "300", "debt_ratio": 55},
            "report_chapters": [{"title": "财务", "analysis": ["ok"]}],
        },
        "legal": {"recommendation": "无异常"},
        "industry": {"recommendation": "中性"},
    }
    report = build_full_due_diligence_report(
        enterprise_name="测试企业",
        sub_reports=sub_reports,
        report_mode="financial_enhanced_dd",
        template=tpl,
    )
    assert "template_sections" in report
    assert report["template"]["name"] == "集成测试"
    # 模板章节顺序保留
    assert [s["title"] for s in report["template_sections"]] == ["企业概况", "财务分析", "行业环境"]
    # 指标经 key_metrics 绑定
    fin = report["template_sections"][1]
    ind = next(c for c in fin["content"] if c["type"] == "indicator" and c["label"] == "营业收入")
    assert ind["value"] == "1200"


def _minimal_research_state():
    return {
        "enterprise_name": "测试企业",
        "objective": "完整贷前尽调",
        "tasks": [],
        "claims": [],
        "gaps": [],
        "evidence": [],
        "planner": {},
        "parsed_intent": {},
        "timeline": [],
        "research_rounds": [],
        "follow_up_tasks": [],
        "tool_traces": {},
    }


def test_deepresearch_report_attaches_template_sections():
    """非功能需求①：DeepResearch 综合报告也应按用户上传模板组织章节。"""
    tpl = parse_markdown(SAMPLE_MD, name="deepresearch模板")
    state = _minimal_research_state()
    from app.template.store import save_template, set_active

    saved = save_template(tpl)
    set_active(saved.id)
    try:
        report = synthesize_research_report(state, template=tpl)
        assert "template_sections" in report
        assert report["template"]["name"] == "deepresearch模板"
        # 模板章节顺序保留
        assert [s["title"] for s in report["template_sections"]] == ["企业概况", "财务分析", "行业环境"]
        fin_sec = report["template_sections"][1]
        # 子报告嵌入：financial 章节（DeepResearch 财务章标题含「财务状况」）
        sub = next(c for c in fin_sec["content"] if c["type"] == "subreport" and c["dimension"] == "financial")
        assert any("财务状况" in ch.get("title", "") for ch in sub["chapters"])
        # 解读位置：取对应维度专项结论（非默认的「AI 解读位置」占位）
        interp = next(c for c in fin_sec["content"] if c["type"] == "interpretation")
        assert "AI 解读位置" not in interp["text"] and interp["text"].strip()
        # 指标位在无财务证据时显式标待补充
        ind = next(c for c in fin_sec["content"] if c["type"] == "indicator" and c["label"] == "营业收入")
        assert ind["value"] == "待补充"
    finally:
        from app.template.store import delete_template

        delete_template(saved.id)
