# backend/app/engines/rebecca/report_generator.py
"""尽调报告生成器"""
from pathlib import Path
from typing import Dict, Any, Optional
from datetime import datetime
import json


# 财务分析报告的维度渲染顺序（key 对应 analyzers.generate_full_analysis 的输出 key）
# P2.2/P2.3 新增的盈利趋势与归因、盈利质量在此接入报告正文。
_FINANCIAL_SECTIONS = [
    ("profitability", "盈利能力"),
    ("solvency", "偿债能力"),
    ("efficiency", "营运能力"),
    ("growth", "成长能力"),
    ("cash_flow", "现金流"),
    ("profitability_trend", "盈利趋势与归因"),
    ("earnings_quality", "盈利质量"),
]


class ReportGenerator:
    """尽调报告生成器"""

    def __init__(self, output_dir: str = "./output/reports"):
        """
        初始化报告生成器

        Args:
            output_dir: 报告输出目录
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(
        self,
        company_name: str,
        analysis_result: Dict[str, Any],
        format: str = "docx",
    ) -> str:
        """
        生成尽调报告

        Args:
            company_name: 公司名称
            analysis_result: 分析结果
            format: 报告格式（docx 或 pdf）

        Returns:
            str: 报告文件路径
        """
        if format == "docx":
            return self._generate_docx(company_name, analysis_result)
        elif format == "pdf":
            return self._generate_pdf(company_name, analysis_result)
        else:
            raise ValueError(f"不支持的报告格式: {format}")

    def _generate_docx(
        self, company_name: str, analysis_result: Dict[str, Any]
    ) -> str:
        """生成 Word 报告"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # 设置标题
            title = doc.add_heading(f"{company_name} 尽职调查报告", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 添加报告信息
            doc.add_paragraph(f"报告生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph("")

            # 添加目录
            doc.add_heading("目录", level=1)
            risk_num = 3 + len(_FINANCIAL_SECTIONS)
            conc_num = risk_num + 1
            toc_items = ["1. 企业基本信息", "2. 财务分析概览"]
            toc_items += [f"{i}. {name}分析" for i, (_, name) in enumerate(_FINANCIAL_SECTIONS, start=3)]
            toc_items += [f"{risk_num}. 风险提示", f"{conc_num}. 结论与建议"]
            for item in toc_items:
                doc.add_paragraph(item, style="List Number")

            doc.add_page_break()

            # 1. 企业基本信息
            doc.add_heading("1. 企业基本信息", level=1)
            doc.add_paragraph(f"企业名称: {company_name}")
            doc.add_paragraph(f"报告类型: 财务尽职调查")
            doc.add_paragraph(f"分析日期: {datetime.now().strftime('%Y-%m-%d')}")

            # 2. 财务分析概览
            doc.add_heading("2. 财务分析概览", level=1)
            doc.add_paragraph(
                f"本报告基于企业提供的财务数据，从以下 {len(_FINANCIAL_SECTIONS)} 个维度进行全面分析:")
            for _, name in _FINANCIAL_SECTIONS:
                doc.add_paragraph(f"• {name}", style="List Bullet")

            # 3~N. 各维度分析
            section_num = 3
            for dim_key, dim_name in _FINANCIAL_SECTIONS:
                doc.add_heading(f"{section_num}. {dim_name}分析", level=1)

                if dim_key in analysis_result:
                    tables = analysis_result[dim_key]
                    for table_name, table_data in tables.items():
                        doc.add_heading(f"{dim_name} - {table_name}", level=2)
                        self._render_table_docx(doc, table_data)
                else:
                    doc.add_paragraph("暂无数据")

                section_num += 1

            # 风险提示
            doc.add_heading(f"{risk_num}. 风险提示", level=1)
            if "risk_assessment" in analysis_result:
                risk_data = analysis_result["risk_assessment"]
                if "risk_summary" in risk_data and not risk_data["risk_summary"].empty:
                    self._render_table_docx(doc, risk_data["risk_summary"])
                else:
                    doc.add_paragraph("暂无重大风险提示")
            else:
                doc.add_paragraph("暂无重大风险提示")

            # 结论与建议
            doc.add_heading(f"{conc_num}. 结论与建议", level=1)
            doc.add_paragraph("基于以上分析，对企业财务状况的总体评价和建议如下:")
            doc.add_paragraph("1. 企业整体财务状况良好，盈利能力稳定")
            doc.add_paragraph("2. 偿债能力处于行业正常水平")
            doc.add_paragraph("3. 建议关注应收账款管理和成本控制")

            # 保存文件
            filename = f"{company_name}_尽调报告_{datetime.now().strftime('%Y%m%d')}.docx"
            filepath = self.output_dir / filename
            doc.save(str(filepath))

            return str(filepath)

        except ImportError:
            raise ImportError("生成 Word 报告需要安装 python-docx: pip install python-docx")
        except Exception as e:
            raise RuntimeError(f"Word 报告生成失败: {str(e)}")

    # ──────────────────────────────────────────────────────────────────────────
    # 通用工具
    # ──────────────────────────────────────────────────────────────────────────

    @staticmethod
    def _esc(text: Any) -> str:
        """转义 XML 特殊字符，供 reportlab Paragraph 使用。"""
        return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    @staticmethod
    def _to_serializable(obj: Any) -> Any:
        """把 DataFrame / 嵌套结构转成可 JSON 序列化的对象。"""
        try:
            import pandas as pd
        except Exception:
            pd = None
        if pd is not None and isinstance(obj, pd.DataFrame):
            return obj.to_dict(orient="records")
        if isinstance(obj, (list, tuple)):
            return [ReportGenerator._to_serializable(x) for x in obj]
        if isinstance(obj, dict):
            return {k: ReportGenerator._to_serializable(v) for k, v in obj.items()}
        return obj

    @staticmethod
    def _render_table_docx(doc, table_data: Any) -> None:
        """把 DataFrame 渲染成 Word 表格；非 DataFrame 回退 json 文本。"""
        try:
            import pandas as pd
        except Exception:
            pd = None
        if pd is not None and isinstance(table_data, pd.DataFrame) and not table_data.empty:
            header = [str(c) for c in table_data.columns]
            records = table_data.astype(str).values.tolist()
            t = doc.add_table(rows=len(records) + 1, cols=len(header))
            try:
                t.style = "Light Grid Accent 1"
            except Exception:
                pass
            for c, h in enumerate(header):
                t.cell(0, c).text = str(h)
            for r, row in enumerate(records, 1):
                for c, val in enumerate(row):
                    t.cell(r, c).text = str(val)
        else:
            doc.add_paragraph(json.dumps(table_data, ensure_ascii=False, indent=2, default=str))

    @staticmethod
    def _render_table_pdf(flow: list, table_data: Any, cn_font: str, P_func) -> None:
        """把 DataFrame 渲染成 reportlab 表格追加到 flow；非 DataFrame 回退 Preformatted。"""
        from reportlab.platypus import Table, TableStyle, Preformatted
        from reportlab.lib import colors
        try:
            import pandas as pd
        except Exception:
            pd = None
        if pd is not None and isinstance(table_data, pd.DataFrame) and not table_data.empty:
            header = [str(c) for c in table_data.columns]
            rows = [[str(c) for c in row] for row in table_data.astype(str).values.tolist()]
            data = [header] + rows
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), cn_font),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            flow.append(t)
        else:
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            pre = ParagraphStyle("PreCn", parent=getSampleStyleSheet()["Code"], fontName=cn_font, fontSize=7.5, leading=10)
            flow.append(Preformatted(
                json.dumps(table_data, ensure_ascii=False, indent=2, default=str), pre))

    @staticmethod
    def _risk_summary_rows(analysis_result: Dict[str, Any]) -> list:
        """从 risk_assessment.risk_summary 提取可读风险条目（兼容 DataFrame / list）。"""
        risk_data = analysis_result.get("risk_assessment", {}) or {}
        summary = risk_data.get("risk_summary")
        rows: list = []
        if hasattr(summary, "iterrows"):
            for _, row in summary.iterrows():
                rows.append(f"[{row.get('风险等级', '未知')}] {row.get('风险描述', '')}")
        elif isinstance(summary, list):
            for item in summary:
                if isinstance(item, dict):
                    rows.append(f"[{item.get('风险等级', '未知')}] {item.get('风险描述', '')}")
                else:
                    rows.append(str(item))
        return rows

    @staticmethod
    def _humanize(key: str) -> str:
        if not key:
            return ""
        s = str(key).replace("_", " ").strip()
        return s[:1].upper() + s[1:] if s else s

    # ──────────────────────────────────────────────────────────────────────────
    # PDF 报告（财务 analysis_result 维度结构，镜像 _generate_docx）
    # ──────────────────────────────────────────────────────────────────────────

    def _generate_pdf(self, company_name: str, analysis_result: Dict[str, Any]) -> str:
        """生成 PDF 报告（reportlab，内置 STSong-Light 中文字体，跨平台可用）。"""
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                        ListFlowable, ListItem, Preformatted, PageBreak)
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            cn = "STSong-Light"
        except Exception:
            cn = "Helvetica"

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("TitleCn", parent=styles["Title"], fontName=cn, fontSize=18, leading=24)
        h1 = ParagraphStyle("H1Cn", parent=styles["Heading1"], fontName=cn, fontSize=14, leading=18)
        body = ParagraphStyle("BodyCn", parent=styles["BodyText"], fontName=cn, fontSize=10, leading=14)
        pre = ParagraphStyle("PreCn", parent=styles["Code"], fontName=cn, fontSize=7.5, leading=10)
        bullet = ParagraphStyle("BulletCn", parent=styles["BodyText"], fontName=cn, fontSize=10, leading=14)

        def P(text: Any, style=body):
            return Paragraph(self._esc(text), style)

        def bullets(items):
            return ListFlowable(
                [ListItem(P(t, bullet), leftIndent=12) for t in items],
                bulletType="bullet", start="•",
            )

        flow = [
            P(f"{company_name} 尽职调查报告", title_style),
            P(f"报告生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"),
            Spacer(1, 6),
        ]

        flow.append(P("目录", h1))
        _risk_num = 3 + len(_FINANCIAL_SECTIONS)
        _conc_num = _risk_num + 1
        _toc = ["1. 企业基本信息", "2. 财务分析概览"]
        _toc += [f"{i}. {name}分析" for i, (_, name) in enumerate(_FINANCIAL_SECTIONS, start=3)]
        _toc += [f"{_risk_num}. 风险提示", f"{_conc_num}. 结论与建议"]
        for item in _toc:
            flow.append(P(item, body))
        flow.append(PageBreak())

        flow.append(P("1. 企业基本信息", h1))
        flow.append(P(f"企业名称: {company_name}"))
        flow.append(P("报告类型: 财务尽职调查"))
        flow.append(P(f"分析日期: {datetime.now().strftime('%Y-%m-%d')}"))

        risk_num = 3 + len(_FINANCIAL_SECTIONS)
        conc_num = risk_num + 1

        flow.append(P("2. 财务分析概览", h1))
        flow.append(P(f"本报告基于企业提供的财务数据，从以下 {len(_FINANCIAL_SECTIONS)} 个维度进行全面分析:"))
        flow.append(bullets([name for _, name in _FINANCIAL_SECTIONS]))

        num = 3
        for dim_key, dim_name in _FINANCIAL_SECTIONS:
            flow.append(P(f"{num}. {dim_name}分析", h1))
            if dim_key in analysis_result:
                tables = analysis_result[dim_key]
                for table_name, table_data in tables.items():
                    flow.append(P(f"{dim_name} - {table_name}", body))
                    self._render_table_pdf(flow, table_data, cn, P)
            else:
                flow.append(P("暂无数据"))
            num += 1

        flow.append(P(f"{risk_num}. 风险提示", h1))
        risk_summary = analysis_result.get("risk_assessment", {}).get("risk_summary")
        if risk_summary is not None and hasattr(risk_summary, "iterrows") and not risk_summary.empty:
            self._render_table_pdf(flow, risk_summary, cn, P)
        else:
            flow.append(P("暂无重大风险提示"))

        flow.append(P(f"{conc_num}. 结论与建议", h1))
        flow.append(P("基于以上分析，对企业财务状况的总体评价和建议如下:"))
        flow.append(bullets(["企业整体财务状况良好，盈利能力稳定",
                             "偿债能力处于行业正常水平",
                             "建议关注应收账款管理和成本控制"]))

        filename = f"{company_name}_尽调报告_{datetime.now().strftime('%Y%m%d')}.pdf"
        filepath = self.output_dir / filename
        doc = SimpleDocTemplate(str(filepath), pagesize=A4, title=f"{company_name} 尽职调查报告")
        doc.build(flow)
        return str(filepath)

    # ──────────────────────────────────────────────────────────────────────────
    # 通用结构化报告渲染（直接渲染 task["report"] 嵌套 dict → docx / pdf）
    # ──────────────────────────────────────────────────────────────────────────

    def generate_from_report(self, report: Dict[str, Any], company_name: str, format: str = "docx") -> str:
        """把结构化授信报告（嵌套 dict）渲染为 docx / pdf / markdown 文件并返回路径。"""
        if format == "docx":
            return self._render_structured_docx(company_name, report)
        elif format == "pdf":
            return self._render_structured_pdf(company_name, report)
        elif format == "md":
            return self._render_structured_md(company_name, report)
        else:
            raise ValueError(f"不支持的报告格式: {format}")

    def _collect_blocks(self, name: str, node: Any, blocks: list, level: int = 2) -> None:
        """把嵌套 dict/list 规整为渲染块：(标题级别 | p | bullets | table, 内容)。"""
        if node is None:
            return
        if isinstance(node, dict):
            if name:
                blocks.append((f"h{min(level, 4)}", self._humanize(name)))
            for k, v in node.items():
                if isinstance(v, (dict, list)) and v:
                    self._collect_blocks(k, v, blocks, level + 1)
                else:
                    blocks.append(("p", f"{self._humanize(k)}：{v}"))
        elif isinstance(node, list):
            if name:
                blocks.append((f"h{min(level, 4)}", self._humanize(name)))
            if node and all(isinstance(i, dict) for i in node):
                keys: list = []
                for i in node:
                    for kk in i.keys():
                        if kk not in keys:
                            keys.append(kk)
                rows = [[str(i.get(kk, "")) for kk in keys] for i in node]
                blocks.append(("table", [keys] + rows))
            else:
                items = [str(i) for i in node if i is not None]
                if items:
                    blocks.append(("bullets", items))
        else:
            text = str(node)
            if name:
                blocks.append(("p", f"{self._humanize(name)}：{text}"))
            else:
                blocks.append(("p", text))

    def _render_structured_docx(self, company_name: str, report: Dict[str, Any]) -> str:
        from docx import Document
        from docx.shared import Pt  # noqa: F401
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        title = doc.add_heading(f"{company_name} 尽职调查报告", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"报告生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        blocks: list = []
        self._collect_blocks("", report, blocks, level=2)
        for kind, content in blocks:
            if kind.startswith("h"):
                doc.add_heading(content, level=int(kind[1]))
            elif kind == "p":
                doc.add_paragraph(str(content))
            elif kind == "bullets":
                for it in content:
                    doc.add_paragraph(str(it), style="List Bullet")
            elif kind == "table":
                if content:
                    t = doc.add_table(rows=len(content), cols=len(content[0]))
                    try:
                        t.style = "Light Grid Accent 1"
                    except Exception:
                        pass
                    for r, row in enumerate(content):
                        for c, val in enumerate(row):
                            t.cell(r, c).text = str(val)

        filename = f"{company_name}_授信报告_{datetime.now().strftime('%Y%m%d')}.docx"
        filepath = self.output_dir / filename
        doc.save(str(filepath))
        return str(filepath)

    def _render_structured_pdf(self, company_name: str, report: Dict[str, Any]) -> str:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                        ListFlowable, ListItem, Table, TableStyle)
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            cn = "STSong-Light"
        except Exception:
            cn = "Helvetica"

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("TitleCn", parent=styles["Title"], fontName=cn, fontSize=18, leading=24)
        h1 = ParagraphStyle("H1Cn", parent=styles["Heading1"], fontName=cn, fontSize=14, leading=18)
        h2 = ParagraphStyle("H2Cn", parent=styles["Heading2"], fontName=cn, fontSize=12, leading=16)
        h3 = ParagraphStyle("H3Cn", parent=styles["Heading3"], fontName=cn, fontSize=11, leading=15)
        body = ParagraphStyle("BodyCn", parent=styles["BodyText"], fontName=cn, fontSize=10, leading=14)

        def P(t: Any, s=body):
            return Paragraph(self._esc(t), s)

        blocks: list = []
        self._collect_blocks("", report, blocks, level=2)
        flow = [
            P(f"{company_name} 尽职调查报告", title_style),
            P(f"报告生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"),
            Spacer(1, 6),
        ]
        for kind, content in blocks:
            if kind == "h1":
                flow.append(P(content, h1))
            elif kind in ("h2", "h3", "h4"):
                flow.append(P(content, h3))
            elif kind == "p":
                flow.append(P(content))
            elif kind == "bullets":
                flow.append(ListFlowable(
                    [ListItem(P(t, body), leftIndent=12) for t in content],
                    bulletType="bullet", start="•"))
            elif kind == "table":
                rows = [[self._esc(c) for c in row] for row in content]
                if rows:
                    t = Table(rows, repeatRows=1)
                    t.setStyle(TableStyle([
                        ("FONTNAME", (0, 0), (-1, -1), cn),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]))
                    flow.append(t)

        filename = f"{company_name}_授信报告_{datetime.now().strftime('%Y%m%d')}.pdf"
        filepath = self.output_dir / filename
        doc = SimpleDocTemplate(str(filepath), pagesize=A4, title=f"{company_name} 尽职调查报告")
        doc.build(flow)
        return str(filepath)

    # ──────────────────────────────────────────────────────────────────────────
    # Markdown 报告（轻量、可版本控制、便于人工复核）
    # ──────────────────────────────────────────────────────────────────────────

    @staticmethod
    def _md_escape_cell(text: str) -> str:
        """转义 Markdown 表格单元格中的管道符与换行。"""
        return str(text).replace("|", "\\|").replace("\n", " ").replace("\r", "")

    def _render_structured_md(self, company_name: str, report: Dict[str, Any]) -> str:
        """把结构化授信报告渲染为 Markdown 文件并返回路径。"""
        blocks: list = []
        self._collect_blocks("", report, blocks, level=2)

        lines: list[str] = []
        lines.append(f"# {company_name} 尽职调查报告")
        lines.append("")
        lines.append(f"> 报告生成日期：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        lines.append("")

        for kind, content in blocks:
            if kind.startswith("h"):
                level = int(kind[1])
                lines.append(f"{'#' * level} {content}")
                lines.append("")
            elif kind == "p":
                lines.append(str(content))
                lines.append("")
            elif kind == "bullets":
                for item in content:
                    lines.append(f"- {item}")
                lines.append("")
            elif kind == "table":
                if content:
                    header = [self._md_escape_cell(cell) for cell in content[0]]
                    lines.append("| " + " | ".join(header) + " |")
                    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for row in content[1:]:
                        lines.append("| " + " | ".join(self._md_escape_cell(cell) for cell in row) + " |")
                    lines.append("")

        filename = f"{company_name}_授信报告_{datetime.now().strftime('%Y%m%d')}.md"
        filepath = self.output_dir / filename
        filepath.write_text("\n".join(lines), encoding="utf-8")
        return str(filepath)
